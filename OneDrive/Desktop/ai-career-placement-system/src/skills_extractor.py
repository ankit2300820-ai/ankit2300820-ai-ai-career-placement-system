"""Resume parsing (PDF/DOCX/TXT) + rule-based skill extraction using a skill taxonomy."""
import io, re

ALIASES = {
    "Python": ["python"], "Java": ["java"], "C++": ["c++"], "C": ["c programming", "c language"],
    "JavaScript": ["javascript", "js", "es6"], "TypeScript": ["typescript"], "SQL": ["sql", "mysql", "postgresql", "sqlite"],
    "R": ["r programming"], "HTML/CSS": ["html", "css", "html5", "css3"], "React": ["react", "reactjs", "react.js"],
    "Node.js": ["node.js", "nodejs", "node js"], "Django": ["django"], "Flask": ["flask"], "Spring Boot": ["spring boot", "springboot"],
    "REST APIs": ["rest api", "rest apis", "restful"], "MongoDB": ["mongodb"], "Git": ["git", "github", "gitlab"],
    "Docker": ["docker"], "Kubernetes": ["kubernetes", "k8s"], "AWS": ["aws", "amazon web services", "ec2", "s3"],
    "Linux": ["linux", "ubuntu", "bash", "shell scripting"], "CI/CD": ["ci/cd", "jenkins", "github actions"],
    "Pandas": ["pandas"], "NumPy": ["numpy"], "Scikit-learn": ["scikit-learn", "sklearn"],
    "Machine Learning": ["machine learning", "ml"], "Deep Learning": ["deep learning", "neural network", "neural networks", "cnn"],
    "TensorFlow": ["tensorflow", "keras", "pytorch"], "NLP": ["nlp", "natural language processing"],
    "Statistics": ["statistics", "statistical", "hypothesis testing"], "Excel": ["excel", "ms excel"],
    "Power BI": ["power bi", "powerbi"], "Tableau": ["tableau"], "Data Structures": ["data structures", "dsa"],
    "Algorithms": ["algorithms", "dsa"], "OOP": ["oop", "oops", "object oriented", "object-oriented"],
    "DBMS": ["dbms", "database management"], "Operating Systems": ["operating system", "operating systems"],
    "Computer Networks": ["computer networks", "networking", "tcp/ip"], "Problem Solving": ["problem solving", "competitive programming", "leetcode"],
    "Cybersecurity": ["cybersecurity", "cyber security", "network security", "penetration testing"], "Selenium": ["selenium", "test automation"],
}
ALL_SKILLS = sorted(ALIASES)
_PATTERNS = {s: re.compile(r"(?<![A-Za-z0-9+#.])(?:%s)(?![A-Za-z0-9+#])" % "|".join(re.escape(a) for a in al), re.I)
             for s, al in ALIASES.items()}

def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(data))
        return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    return data.decode("utf-8", errors="ignore")

def extract_skills(text: str):
    return sorted(s for s, p in _PATTERNS.items() if p.search(text))

def extract_meta(text: str):
    """Pull simple facts (email, CGPA, project/internship mentions) from resume text."""
    meta = {}
    if m := re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text): meta["email"] = m.group(0)
    if m := re.search(r"c\.?g\.?p\.?a\.?\s*[:\-]?\s*(\d{1,2}(?:\.\d+)?)", text, re.I):
        v = float(m.group(1)); meta["cgpa"] = v if v <= 10 else None
    meta["internship_mentions"] = len(re.findall(r"\bintern(ship)?\b", text, re.I))
    meta["project_mentions"] = len(re.findall(r"\bprojects?\b", text, re.I))
    return meta
